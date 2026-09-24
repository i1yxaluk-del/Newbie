#!/usr/bin/env python3
"""Собирает единый договор MSPShield из Markdown в DOCX.
Назначение: воспроизводимая версия без ручного копирования.
Где запускать: из contracts после проверки плейсхолдеров.
Побочные эффекты: перезаписывает MSP_SERVICE_AGREEMENT.docx.
Проверка: открыть DOCX/PDF, проверить таблицы и разрывы страниц.
Откат: исходный Markdown остаётся неизменным.
"""
from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
BASE=Path(__file__).resolve().parent
SOURCE=BASE/'MSP_SERVICE_AGREEMENT.md'; OUTPUT=BASE/'MSP_SERVICE_AGREEMENT.docx'
def shade(cell,fill):
 p=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); p.append(e)
def rich(doc,text,style=None):
 p=doc.add_paragraph(style=style)
 for part in re.split(r'(\*\*.*?\*\*)',text):
  r=p.add_run(part[2:-2] if part.startswith('**') else part.replace('`','')); r.bold=part.startswith('**')
 return p
def build():
 doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2.2); sec.right_margin=Cm(1.8)
 doc.styles['Normal'].font.name='Arial'; doc.styles['Normal'].font.size=Pt(10.5)
 for n,s,c in [('Title',18,'1F4E79'),('Heading 1',14,'1F4E79')]:
  doc.styles[n].font.name='Arial'; doc.styles[n].font.size=Pt(s); doc.styles[n].font.color.rgb=RGBColor.from_string(c)
 lines=SOURCE.read_text(encoding='utf-8').splitlines(); i=0
 while i<len(lines):
  line=lines[i].rstrip()
  if not line: i+=1; continue
  if line.startswith('# '):
   if doc.paragraphs: doc.add_page_break()
   p=doc.add_paragraph(line[2:],style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  elif line.startswith('## '): doc.add_paragraph(line[3:],style='Heading 1')
  elif line.startswith('> '): rich(doc,line[2:]).paragraph_format.left_indent=Cm(.7)
  elif line.startswith('- '): rich(doc,line[2:],'List Bullet')
  elif line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r':?-{3,}:?',c or '') for c in cells): rows.append(cells)
    i+=1
   t=doc.add_table(rows=len(rows),cols=max(map(len,rows))); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
   for ri,row in enumerate(rows):
    for ci,val in enumerate(row):
     t.cell(ri,ci).text=val
     for r in t.cell(ri,ci).paragraphs[0].runs: r.font.name='Arial'; r.font.size=Pt(8.5); r.bold=ri==0
     if ri==0: shade(t.cell(ri,ci),'DCE6F1')
   continue
  else: rich(doc,line)
  i+=1
 for s in doc.sections:
  p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('MSPShield · единый шаблон · требуется юридическая проверка').font.size=Pt(8)
 doc.core_properties.title='Договор оказания управляемых IT-услуг MSPShield'; doc.save(OUTPUT); print(OUTPUT)
if __name__=='__main__': build()
